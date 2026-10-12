# -*- coding: utf-8 -*-
"""
ShelfMind Technical Project Report Generator
Builds a publication-grade, academic technical report in Word (.docx) format.
Enforces:
- Times New Roman throughout
- Headings at 14pt bold
- Body text at 11pt, justified
- Professional Table of Contents with page references
- Technical formulas (MDP, Bellman, Laplacian, Q-learning, DQN)
- Embedded screenshots and analytical diagrams
- Detailed technical tables
- Peer-reviewed research references
- Page target <= 25 pages
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="1F497D"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_footer_page_number(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def add_footer_total_pages(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'NUMPAGES')
    run._r.append(fldSimple)

def style_paragraph(p, font_name="Times New Roman", font_size=11, bold=False, italic=False, color_rgb=(0,0,0), alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.15):
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

def add_body_p(doc, text, bold_prefix="", italic_note=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4.5)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)

    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(11)
    r_body.font.bold = False
    r_body.font.color.rgb = RGBColor(25, 25, 25)

    if italic_note:
        r_it = p.add_run(italic_note)
        r_it.font.name = "Times New Roman"
        r_it.font.size = Pt(11)
        r_it.font.italic = True
        r_it.font.color.rgb = RGBColor(60, 60, 60)

    return p

def add_heading_14(doc, title, space_before=14, space_after=5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(16, 44, 87) # Deep academic navy
    return p

def add_subheading_12(doc, title, space_before=10, space_after=3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 78, 120)
    return p

def add_formula_block(doc, formula_text, eq_number=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.4)

    r = p.add_run(formula_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.italic = True
    r.font.bold = True
    r.font.color.rgb = RGBColor(20, 20, 20)

    if eq_number:
        r_num = p.add_run(f"    \t({eq_number})")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(10.5)
        r_num.font.italic = False
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(80, 80, 80)
    return p

def add_figure(doc, img_path, caption, width=Inches(5.8)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        p_cap.paragraph_format.keep_with_next = False
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(60, 60, 60)

print("Helper definitions loaded successfully.")
